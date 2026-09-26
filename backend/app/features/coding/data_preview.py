"""A first look at an uploaded data file: its columns, their types, how many
rows, and the first few — the check you make before asking for an analysis.

Shown to the user as a table at upload, and summarised into the coding
prompt so the model writes code against the real column names instead of
guessing them.

Stdlib csv/json plus openpyxl and pyarrow (both direct dependencies), never
pandas, which is only a transitive one here. Best effort: a file that can't
be read as a table gets no preview and the upload goes on as before.
"""
from __future__ import annotations

import csv
import io
import json
import logging
import math
import re
from dataclasses import dataclass
from datetime import date, datetime, time
from typing import Any

logger = logging.getLogger(__name__)

PREVIEW_ROWS = 5
MAX_COLUMNS = 30
MAX_CELL_CHARS = 80
# Rows looked at to guess a column's type. Enough to get past a few blanks
# at the top without reading the whole file twice.
_TYPE_SAMPLE = 200

_MISSING = {"", "na", "n/a", "nan", "null", "none", "-"}
_INT = re.compile(r"[+-]?\d+")
_FLOAT = re.compile(r"[+-]?(\d+\.\d*|\.\d+|\d+)([eE][+-]?\d+)?")
_DATE = re.compile(r"\d{4}-\d{2}-\d{2}")
_DATETIME = re.compile(r"\d{4}-\d{2}-\d{2}[ T]\d{2}:\d{2}(:\d{2}(\.\d+)?)?(Z|[+-]\d{2}:?\d{2})?")



@dataclass
class _Parsed:
    header: list[str]
    sample: list[list[Any]]  # the first rows, up to _TYPE_SAMPLE
    total_rows: int
    types: list[str] | None = None  # when the format stores them (parquet)
    sheet: str | None = None


def table_preview(suffix: str, content: bytes) -> dict | None:
    """Columns with a guessed type, the first rows, and the totals — or None
    when `suffix` isn't a tabular format or the file can't be read as one."""
    reader = _READERS.get(suffix)
    if reader is None:
        return None
    try:
        parsed = reader(content)
    except Exception as exc:  # noqa: BLE001 — a bad file must not fail the upload
        logger.info("coding.preview_unavailable suffix=%s error=%s", suffix, type(exc).__name__)
        return None
    if parsed is None or not parsed.header:
        return None

    width = len(parsed.header)
    sample = [_fit(row, width) for row in parsed.sample]
    columns = [
        {"name": name, "type": parsed.types[i] if parsed.types else _column_type([row[i] for row in sample])}
        for i, name in enumerate(parsed.header[:MAX_COLUMNS])
    ]
    table = {
        "columns": columns,
        "rows": [[_cell(value) for value in row[:MAX_COLUMNS]] for row in sample[:PREVIEW_ROWS]],
        "total_rows": parsed.total_rows,
        "total_columns": width,
    }
    if parsed.sheet:
        table["sheet"] = parsed.sheet
    return table


def describe_table(table: Any) -> str:
    """One line for the coding prompt: "1,204 rows × 3 columns: id (int), …".
    `table` comes back from the browser with the request, so nothing about
    its shape is trusted."""
    try:
        columns = [
            f"{str(c.get('name', ''))[:60]} ({str(c.get('type', ''))[:12]})"
            for c in table.get("columns", [])[:MAX_COLUMNS]
            if isinstance(c, dict)
        ]
        rows = int(table.get("total_rows", 0))
        width = int(table.get("total_columns", len(columns)))
    except (AttributeError, TypeError, ValueError):
        return ""
    if not columns:
        return ""
    more = f", … (+{width - len(columns)} more)" if width > len(columns) else ""
    return f"{rows:,} rows × {width} columns: {', '.join(columns)}{more}"


# ── readers ─────────────────────────────────────────────────────────────


def _decode(content: bytes) -> str:
    # utf-8-sig: Excel's "CSV UTF-8" starts with a BOM that would otherwise
    # end up glued to the first column's name.
    return content.decode("utf-8-sig", errors="replace")


def _read_delimited(content: bytes, delimiter: str | None) -> _Parsed | None:
    text = _decode(content)
    if delimiter is None:
        try:
            delimiter = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|").delimiter
        except csv.Error:
            delimiter = ","
    rows = csv.reader(io.StringIO(text), delimiter=delimiter)
    header = next(rows, None)
    if not header:
        return None
    sample: list[list[Any]] = []
    total = 0
    for row in rows:
        if not any(cell.strip() for cell in row):
            continue
        total += 1
        if len(sample) < _TYPE_SAMPLE:
            sample.append(row)
    return _Parsed([h.strip() for h in header], sample, total)


def _records(records: list[dict]) -> _Parsed:
    header: list[str] = []
    seen: set[str] = set()
    for record in records[:_TYPE_SAMPLE]:
        for key in record:
            if key not in seen:
                seen.add(key)
                header.append(key)
    sample = [[record.get(key) for key in header] for record in records[:_TYPE_SAMPLE]]
    return _Parsed(header, sample, len(records))


def _read_json(content: bytes) -> _Parsed | None:
    data = json.loads(_decode(content))
    if not isinstance(data, list) or not data or not all(isinstance(r, dict) for r in data):
        return None
    return _records(data)


def _read_jsonl(content: bytes) -> _Parsed | None:
    lines = [line for line in _decode(content).splitlines() if line.strip()]
    first = [json.loads(line) for line in lines[:_TYPE_SAMPLE]]
    if not first or not all(isinstance(r, dict) for r in first):
        return None
    parsed = _records(first)
    parsed.total_rows = len(lines)
    return parsed


def _read_xlsx(content: bytes) -> _Parsed | None:
    from openpyxl import load_workbook

    book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        sheet = book.worksheets[0]
        rows = (row for row in sheet.iter_rows(values_only=True) if any(v is not None for v in row))
        header = next(rows, None)
        if header is None:
            return None
        sample: list[list[Any]] = []
        total = 0
        for row in rows:
            total += 1
            if len(sample) < _TYPE_SAMPLE:
                sample.append(list(row))
        names = ["" if h is None else str(h).strip() for h in header]
        label = sheet.title if len(book.worksheets) == 1 else f"{sheet.title} (1/{len(book.worksheets)} sheet)"
        return _Parsed(names, sample, total, sheet=label)
    finally:
        book.close()


def _read_parquet(content: bytes) -> _Parsed:
    import pyarrow.parquet as pq

    file = pq.ParquetFile(io.BytesIO(content))
    schema = file.schema_arrow
    header = list(schema.names)
    first = next(file.iter_batches(batch_size=PREVIEW_ROWS), None)
    sample = [[record.get(name) for name in header] for record in (first.to_pylist() if first else [])]
    types = [_arrow_type(field.type) for field in schema]
    return _Parsed(header, sample, file.metadata.num_rows, types=types)


def _arrow_type(arrow_type) -> str:
    import pyarrow.types as pat

    if pat.is_boolean(arrow_type):
        return "bool"
    if pat.is_integer(arrow_type):
        return "int"
    if pat.is_floating(arrow_type) or pat.is_decimal(arrow_type):
        return "float"
    if pat.is_timestamp(arrow_type):
        return "datetime"
    if pat.is_date(arrow_type):
        return "date"
    if pat.is_string(arrow_type) or pat.is_large_string(arrow_type):
        return "text"
    return str(arrow_type)[:12]


_READERS = {
    ".csv": lambda content: _read_delimited(content, None),
    ".tsv": lambda content: _read_delimited(content, "\t"),
    ".json": _read_json,
    ".jsonl": _read_jsonl,
    ".xlsx": _read_xlsx,
    ".parquet": _read_parquet,
}


# ── cells and types ─────────────────────────────────────────────────────


def _fit(row: list[Any], width: int) -> list[Any]:
    """Ragged rows (common in hand-made CSVs) padded or cut to the header."""
    return list(row[:width]) + [None] * (width - len(row))


def _kind(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return "bool"
    if isinstance(value, int):
        return "int"
    if isinstance(value, float):
        return None if math.isnan(value) else "float"
    if isinstance(value, datetime):
        return "datetime"
    if isinstance(value, date):
        return "date"
    if isinstance(value, time):
        return "time"
    if isinstance(value, (dict, list)):
        return "json"
    text = str(value).strip()
    if text.lower() in _MISSING:
        return None
    if _INT.fullmatch(text):
        return "int"
    if _FLOAT.fullmatch(text):
        return "float"
    if text.lower() in ("true", "false"):
        return "bool"
    if _DATE.fullmatch(text):
        return "date"
    if _DATETIME.fullmatch(text):
        return "datetime"
    return "text"


def _column_type(values: list[Any]) -> str:
    kinds = {kind for value in values if (kind := _kind(value))}
    if not kinds:
        return "empty"
    if kinds <= {"int", "float"}:
        return "float" if "float" in kinds else "int"
    if kinds <= {"date", "datetime"}:
        return "datetime" if "datetime" in kinds else "date"
    return kinds.pop() if len(kinds) == 1 else "text"


def _cell(value: Any) -> Any:
    """A JSON-safe, display-sized version of one value."""
    if value is None or isinstance(value, (bool, int)):
        return value
    if isinstance(value, float):
        return value if math.isfinite(value) else None
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    text = json.dumps(value, ensure_ascii=False, default=str) if isinstance(value, (dict, list)) else str(value)
    return text if len(text) <= MAX_CELL_CHARS else text[: MAX_CELL_CHARS - 1] + "…"
